"""Fichas de palanca por empresa y fichas de intervención (negocio/software y agentes) para el consolidado de P2R.

Siguen los formularios de la aplicación: «Palanca · Definición» (sin el compromiso de valor, que va con la
cuantificación) y «Nueva intervención». Las fechas salen de las propuestas equivalentes; las horas de personas, de la
dedicación cargada (solo hay en Ondupet).
"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font

import generar_intervenciones_mgt as mgt
import generar_intervenciones_ondupack as opack
import generar_intervenciones_ondupet as opet
from generar_intervenciones_ondupet import D, F, cabecera, celda, fill, titulo
from hojas_empresa import HALLAZGOS, INTERVENCIONES, PALANCAS, descripcion

EMPRESAS = ('Ondupet', 'MGT', 'Ondupack Almendralejo')
FECHAS = {emp: {it['nombre']: (it['ini'], it['fin'], it) for it in modulo.INTERVENCIONES}
          for emp, modulo in (('Ondupet', opet), ('Ondupack Almendralejo', opack), ('MGT', mgt))}
IT = 'Sistemas / IT (dar de alta en el censo)'
OWNER = 'el owner'

# ---------------------------------------------------------------- palancas
LINEA = {'Circulante': 'Circulante', 'Fuga de P&L': 'P&L · fuga', 'Transformación': 'Transformación (capacidad de línea)',
         '—': 'Oportunidad (ingreso)'}
QUITAR = {('Ondupet', 'P2R-03', 'OPP-02'), ('MGT', 'P2R-03', 'OPP-02'), ('Ondupet', 'P2R-01', 'INF-16')}
FAMILIA_CORREGIDA = {('MGT', 'INF-01'): 'Circulante'}

# (empresa, palanca): owner, equipo de apoyo, horas de personas, contador, qué cambiar en la aplicación
FICHA_PALANCA = {
    ('Ondupet', 'P2R-01'): (
        'Manuel Jesús Trinidad Montero', 'Victoriano Pérez Carrasco · Monica Gordillo Guerrero',
        'Previsión y Precisión · Victoriano Pérez · 86 h/año (5 % en Demanda y previsión); INT-01 le quita ≈ 26 h',
        'Precisión y sesgo por familia: previsión cerrada frente a expediciones, en GIO',
        'Quitar la línea de INF-16 (ya no está vivo) y añadir «+ Horas de personas» en Previsión y Precisión para que cuente '
        'la capacidad de INT-01 (903 €).'),
    ('Ondupet', 'P2R-02'): (
        'Victoriano Pérez Carrasco', 'Manuel Jesús Trinidad Montero · ' + IT,
        'Plan Maestro · Manuel Jesús Trinidad 258 h/año y Monica Gordillo 344 h/año (S&OP y Plan Maestro)',
        '% de saturación por línea de extrusión y termoformado, en GIO',
        'Añadir «+ Horas de personas» en Plan Maestro (INT-02). Revisar la dedicación de Victoriano: no tiene % en S&OP y Plan '
        'Maestro aunque hace el plan de fabricación y los temporales.'),
    ('Ondupet', 'P2R-03'): (
        'Monica Gordillo Guerrero', 'Victoriano Pérez Carrasco · Manuel Jesús Trinidad Montero',
        '— (ninguna intervención cambia cajas)', 'Cobertura en días de producto terminado por segmento, en GIO',
        'Quitar OPP-02 de «qué ataca»: lo cura P2R-04.'),
    ('Ondupet', 'P2R-04'): (
        'Manuel Jesús Trinidad Montero', 'Victoriano Pérez Carrasco · Monica Gordillo Guerrero',
        'Plan Maestro (Cartera Pedidos) · Manuel Jesús Trinidad 258 h/año y Monica Gordillo 344 h/año',
        '% de pedidos entregados en la primera fecha confirmada: fecha confirmada frente a fecha de carga, en GIO',
        'Añadir «+ Horas de personas» en Plan Maestro (INT-03).'),
    ('Ondupet', 'P2R-06'): (
        'Victoriano Pérez Carrasco', 'Producción (dar de alta en el censo) · ' + IT,
        'Programación · Victoriano Pérez · 430 h/año (25 %)', 'Horas de cambio de molde y altura por semana, en los partes de producción',
        'Añadir «+ Horas de personas» en Programación (INT-04).'),
    ('Ondupet', 'P2R-08'): (
        'Victoriano Pérez Carrasco', 'Manuel Jesús Trinidad Montero',
        '— (ninguna intervención cambia cajas; si ACC-11 marca su caja: Seguimiento y control · Victoriano · 258 h/año)',
        'OTIF contra la fecha del cliente: GIO y albaranes de SAP', ''),
    ('Ondupack Almendralejo', 'P2R-01'): (
        'Miguel Ángel Rosado Martínez', 'Maria Teresa Espinosa Paredes · Tomás Buenavida Santiago',
        '— (ninguna intervención cambia cajas)', '% de la demanda repetitiva con previsión, en GIO', ''),
    ('Ondupack Almendralejo', 'P2R-02'): (
        'José María Baquero Jaen', 'Miguel Ángel Rosado Martínez · Karim Karach Boujtat · Tomás Buenavida Santiago',
        'Carga Capacidad · Miguel Ángel Rosado y Tomás Buenavida (sin dedicación cargada: 0 h)',
        '% de saturación de la 924 con la velocidad real, en GIO',
        'Añadir «+ Horas de personas» en Carga Capacidad (INT-01), después de cargar la dedicación de la celda.'),
    ('Ondupack Almendralejo', 'P2R-05'): (
        'Angélica Lemus Fernández', 'Miguel Ángel Rosado Martínez · Tomás Buenavida Santiago',
        'Papel y Auxiliares · Miguel Ángel Rosado y Angélica Lemus (sin dedicación cargada: 0 h)',
        'Órdenes liberadas con material pendiente (GIO) y € de compras urgentes (SAP)',
        'Anclar ACC-07 (agente) a Papel y Auxiliares y añadir «+ Horas de personas» en esa caja.'),
    ('Ondupack Almendralejo', 'P2R-06'): (
        'Miguel Ángel Rosado Martínez', 'José María Baquero Jaen · Tomás Buenavida Santiago · ' + IT,
        'Converting · Miguel Ángel Rosado (sin dedicación cargada: 0 h)',
        'Horas de cambio en onduladora y converting, con los escenarios SMED de GIO',
        'Añadir «+ Horas de personas» en Converting (INT-03).'),
    ('Ondupack Almendralejo', 'P2R-08'): (
        'Maria Teresa Espinosa Paredes', 'Miguel Ángel Rosado Martínez',
        'Seguimiento del Servicio · Miguel Ángel Rosado y Teresa Espinosa (sin dedicación cargada: 0 h)',
        'OTIF contra la fecha propuesta por el CTP, en GIO',
        'Añadir «+ Horas de personas» en Seguimiento del Servicio (INT-02).'),
    ('MGT', 'P2R-01'): (
        'Albert Fontanet Lopez', 'Joan Carles Girones Campaña · Angélica Lemus Fernández (Ondupack)',
        '— (ninguna intervención cambia cajas)', '% de la demanda (Ondupack y clientes repetitivos) con previsión por calidad y ancho',
        ''),
    ('MGT', 'P2R-02'): (
        'Joan Carles Girones Campaña', 'Albert Fontanet Lopez',
        'Carga de Capacidad · Joan Carles Girones (sin dedicación cargada: 0 h)',
        'Horas de máquina a 3 turnos sin cartera, en GIO y en los partes de producción',
        'Añadir «+ Horas de personas» en Carga de Capacidad (INT-01).'),
    ('MGT', 'P2R-03'): (
        'Joan Carles Girones Campaña', 'Angélica Lemus Fernández (Ondupack) · Albert Fontanet Lopez',
        'Plan Fabricación · Joan Carles Girones (sin dedicación cargada: 0 h)',
        '€ y días de cobertura del stock de Ondupack en MGT, en SAP',
        'Quitar OPP-02 de «qué ataca» (lo cura P2R-04) y añadir «+ Horas de personas» en Plan Fabricación (INT-02).'),
    ('MGT', 'P2R-04'): (
        'Albert Fontanet Lopez', 'Joan Carles Girones Campaña',
        'Demanda Externa · Joan Carles Girones, que hace de atención al cliente (sin dedicación cargada: 0 h)',
        '% de pedidos entregados en la fecha pedida, en SAP',
        'Añadir «+ Horas de personas» en Demanda Externa (INT-04) y la intervención «Fecha confirmada por Planificación».'),
    ('MGT', 'P2R-05'): (
        'Marta Boyo Martinez', 'Joan Carles Girones Campaña · ' + IT,
        'MRP · Marta Boyo y Joan Carles Girones (sin dedicación cargada: 0 h)',
        'Cobertura en días de materias primas y auxiliares, en SAP', 'Añadir «+ Horas de personas» en MRP (INT-03).'),
    ('MGT', 'P2R-06'): (
        'Joan Carles Girones Campaña', 'Ramón (algoritmo, dar de alta en el censo)',
        'Programación · Joan Carles Girones (sin dedicación cargada: 0 h)',
        'Horas de cambio de calidad y ancho por semana, en los partes de producción',
        'Añadir «+ Horas de personas» en Programación (INT-05).'),
    ('MGT', 'P2R-07'): (
        'Joan Carles Girones Campaña', 'Marta Boyo Martinez · ' + IT,
        '— (si ACC-11 marca su caja: Liberación de la Orden · Joan Carles Girones)',
        'Órdenes liberadas con comprobación manual o en papel, en GIO',
        'Añadir la intervención «Orden digital a producción»: INF-08 habla de las órdenes en papel.'),
    ('MGT', 'P2R-08'): (
        'Albert Fontanet Lopez', 'Joan Carles Girones Campaña · Angélica Lemus Fernández (Ondupack)',
        'Entregas a Ondupack y Cierre · Joan Carles Girones (sin dedicación cargada: 0 h)',
        'OTIF de clientes externos y de Ondupack, en SAP',
        'Añadir «+ Horas de personas» en Entregas a Ondupack y Cierre (INT-06).'),
}
PROYECTO = {'Ondupet': 'Piloto P2R Ondupet 2027', 'Ondupack Almendralejo': 'Despliegue P2R Ondupack 2027',
            'MGT': 'Despliegue P2R papeleras (MGT y PDA) 2027-28'}

# ---------------------------------------------------------------- intervenciones de negocio y software
# (empresa, código): propuesta de la que salen las fechas (o fechas propias), responsable si no es el owner,
#                    tareas del flujo que conviene marcar, nota
NEG = {
    ('Ondupet', 'ACC-01'): ('Calendario de previsión con día de cierre', OWNER, '', 'Ya rellenado en la aplicación: solo faltan fechas.'),
    ('Ondupet', 'ACC-02'): ('Registro diario de agotados y demanda perdida', 'Victoriano Pérez Carrasco', '', ''),
    ('Ondupet', 'ACC-03'): ('Carga frente a capacidad automática', IT,
                            'Plan Maestro › Plan de Fabricación → Decisión de turnos y balanceo de materiales → se simplifica',
                            'Ya rellenado en la aplicación: solo faltan fechas.'),
    ('Ondupet', 'ACC-04'): ('Revisión mensual de capacidad con dirección de planta', OWNER, '', ''),
    ('Ondupet', 'ACC-05'): ('Temporales automáticos', IT, 'Plan Maestro › Temporales → Reserva de Ciclo Por Familia → la hace el sistema', ''),
    ('Ondupet', 'ACC-06'): ('Criterios de segmentación y estrategia de stock', OWNER, '', ''),
    ('Ondupet', 'ACC-07'): ('Responsable del stock sin movimiento', 'Manuel Jesús Trinidad Montero', '', ''),
    ('Ondupet', 'ACC-08'): ((D(1, 9, 2027), D(30, 11, 2027)), IT, '', ''),
    ('Ondupet', 'ACC-09'): ('Política de fecha confirmada', OWNER, '', ''),
    ('Ondupet', 'ACC-10'): ('Revisión mensual de causas de atraso', OWNER, '', ''),
    ('Ondupet', 'ACC-11'): ('Cuadro de mando de servicio con alertas', IT,
                            'Seguimiento y control › Seguimiento Diario Cumplimiento Servicio → la hace el sistema', ''),
    ('Ondupack Almendralejo', 'ACC-01'): ((D(1, 6, 2027), D(31, 12, 2027)), OWNER, '', ''),
    ('Ondupack Almendralejo', 'ACC-02'): ('Registro de la demanda perdida en GIO', IT,
                                          'Entrada de Pedidos y Fecha › nueva tarea «Registro de la demanda perdida» → la hace el sistema', ''),
    ('Ondupack Almendralejo', 'ACC-03'): ('Revisión mensual de capacidad y cuellos de botella', OWNER, '', ''),
    ('Ondupack Almendralejo', 'ACC-04'): ('Necesidad de papel calculada desde el plan', IT,
                                          'Papel y Auxiliares › Necesidad de papel por consumo pasado → la hace el sistema', ''),
    ('Ondupack Almendralejo', 'ACC-05'): ('Stock de papel fiable sin esperar al inventario', IT, '', ''),
    ('Ondupack Almendralejo', 'ACC-06'): ('Previsión de papel compartida con MGT/PDA', OWNER, '', ''),
    ('Ondupack Almendralejo', 'ACC-08'): ('KPI oficial de servicio', OWNER, '', 'Lo decide Dirección con Comercial.'),
    ('Ondupack Almendralejo', 'ACC-09'): ('Reunión semanal con KPIs y causas', 'Miguel Ángel Rosado Martínez', '', ''),
    ('Ondupack Almendralejo', 'ACC-10'): ('Copiloto de estado de pedidos', IT,
                                          'Seguimiento del Servicio › Seguimiento de fechas atrasadas → se simplifica', ''),
    ('Ondupack Almendralejo', 'ACC-11'): ('Algoritmo de combinación de onduladora', IT,
                                          'Onduladora › Combinación y secuencia de onduladora → se simplifica; '
                                          'Cartera › Agrupar pedidos repetitivos → la hace el sistema', 'Ya está en el plan de sistemas.'),
    ('MGT', 'ACC-01'): ('Previsión de papel de Ondupack por ciclo', 'Joan Carles Girones Campaña (con Angélica, Ondupack)', '', ''),
    ('MGT', 'ACC-02'): ('Previsión de los clientes externos repetitivos', OWNER, '', ''),
    ('MGT', 'ACC-03'): ('Política de llenado de capacidad', OWNER, '', 'Decisión con Dirección.'),
    ('MGT', 'ACC-04'): ('Revisión mensual de capacidad y ciclos', OWNER, '', ''),
    ('MGT', 'ACC-05'): ('Pull system con Ondupack', OWNER, '', 'Se acuerda con Angélica (Ondupack).'),
    ('MGT', 'ACC-06'): ('MRP de materias primas en SAP', IT, 'MRP › MRP de materias primas en Excel → la hace el sistema', ''),
    ('MGT', 'ACC-07'): ('Algoritmo con lo ya fabricado y fecha límite de cartera', 'Ramón (algoritmo, dar de alta en el censo)',
                        'Programación › Agrupar por calidad y ancho → desaparece; Algoritmo de combinación por ciclos → se simplifica',
                        ''),
    ('MGT', 'ACC-08'): ('Pilar de servicio semanal', OWNER, '', ''),
    ('MGT', 'ACC-09'): ('KPI de servicio', OWNER, '', ''),
    ('MGT', 'ACC-10'): ('Visibilidad de stock y cargas para Ondupack', IT,
                        'Entregas a Ondupack y Cierre › Enviar el stock a Ondupack cada día → desaparece', ''),
    ('MGT', 'ACC-11'): ('Disponibilidad de material por orden', IT,
                        'Liberación de la Orden › Comprobar a mano que hay material → la hace el sistema', ''),
}

# ---------------------------------------------------------------- agentes
# (empresa, código): propuesta (fechas, cobertura, resultado, criterio), caja, tareas del as-is y qué les pasa,
#                    problema, cómo funciona hoy, qué queremos cambiar (comité)
AGENTES = {
    ('Ondupet', 'INT-01'): (
        'Agente de análisis de desviaciones', 'Previsión y Precisión',
        'Comprobación de Accuracy Mes en Curso → la ejecuta un agente (toda la tarea) · ¿Buen % de Accuracy? → la ejecuta un '
        'agente · Corrección Previsiones Futuros Meses → la ejecuta un agente (solo una parte: propone, Comercial aprueba)',
        'Seguimiento manual de las posibles desviaciones de previsión del mes en curso; llegamos tarde a la corrección, así que se '
        'fabrica material por encima o por debajo de su necesidad y hay roturas.',
        'Seguimiento diario en GIO de lo expedido frente a la previsión del mes en curso.',
        'El agente compara cada día lo expedido con la previsión, detecta las familias que se desvían y propone la corrección de '
        'los meses siguientes con su causa; Planificación y Comercial la aprueban.'),
    ('Ondupet', 'INT-02'): (
        'Agente de escenarios de capacidad', 'Plan Maestro',
        'Plan de Fabricación → Decisión de turnos y balanceo de materiales → la ejecuta un agente (solo una parte: propone, '
        'Planificación decide)',
        'Los turnos y el balanceo entre extrusión y termoformado se deciden sin cruzar carga y capacidad: entran necesidades que no '
        'caben o quedan líneas ociosas.',
        'El planificador decide turnos y reparto de carga a mano, por experiencia, al preparar el plan de fabricación.',
        'Un agente propone cada mes el reparto de turnos y de carga por línea con su impacto en horas y €; Planificación lo '
        'aprueba en la revisión de capacidad.'),
    ('Ondupet', 'INT-03'): (
        'Propuesta de fecha con temporales y capacidad libre', 'Plan Maestro',
        'Cartera Pedidos → la ejecuta un agente (solo una parte: propone fecha y notas de carga)',
        'Se mantiene la fecha del cliente aunque se sabe que no se llega, la cartera se revisa por separado en Comercial y en '
        'Planificación y las notas de carga se preparan a mano.',
        'Revisión diaria de la cartera en GIO por Comercial y por Planificación; las fechas no se cambian y las cargas se preparan '
        'a mano.',
        'Un agente propone para cada pedido una fecha viable con la capacidad y los temporales y prepara las notas de carga; '
        'Comercial confirma la fecha al cliente por O2C.'),
    ('Ondupet', 'INT-04'): (
        'Agente de urgencias', 'Programación',
        '¿Cambios de última hora? → la ejecuta un agente (calcula el impacto) · Recuenciar urgencias y prioridades → la ejecuta un '
        'agente (solo una parte: propone la nueva secuencia)',
        'Las urgencias llegan por teléfono y se resecuencian a mano, sin zona congelada ni reglas; si no hay acuerdo, decide el '
        'director de planta.',
        'El planificador recibe la llamada, revisa a mano el programa en GIO y resecuencia.',
        'Un agente calcula el impacto de cada urgencia (pedidos que se retrasan y cambios extra), aplica las reglas de prioridad y '
        'propone la decisión; solo los conflictos llegan a dirección.'),
    ('Ondupack Almendralejo', 'ACC-07'): (
        'Agente de faltas de materia prima', 'Papel y Auxiliares',
        'Necesidad de papel por consumo pasado → la ejecuta un agente (solo una parte: revisa las faltas por orden y propone la '
        'acción)',
        'La falta de materia prima es el primer problema de la planta: las faltas se ven tarde y se resuelven a mano entre '
        'Planificación y Compras.',
        'El papel se pide por consumo pasado y GIO solo ve su stock; las faltas aparecen al liberar y se autorizan a mano con '
        'Compras.',
        'Un agente revisa cada día las faltas por orden y propone la acción (adelantar el pedido a MGT, un sustituto aprobado o '
        'mover la orden); Planificación y Compras deciden.'),
    ('Ondupack Almendralejo', 'INT-01'): (
        'Agente de escenarios de capacidad', 'Carga Capacidad',
        'Carga frente a capacidad con escenarios de 2 a 5 turnos → la ejecuta un agente (prepara los escenarios) · Replanificar '
        'ruta, fecha o turnos → la ejecuta un agente (solo una parte: propone, Planificación decide)',
        'La 924 no alcanza los metros planificados y la sobresaturación se ve tarde, por las fechas o porque se llena el almacén '
        'intermedio.',
        'Planificación compara a mano escenarios de 2 a 5 turnos en la carga-capacidad de GIO y replanifica ruta, fecha o turnos '
        'cuando hay saturación.',
        'Un agente simula turnos, rutas alternativas para descargar la 924 y agrupaciones, y propone el plan con su impacto; '
        'Planificación decide.'),
    ('Ondupack Almendralejo', 'INT-02'): (
        'Agente de causa raíz', 'Seguimiento del Servicio',
        'Seguimiento de fechas atrasadas → la ejecuta un agente (solo una parte: clasifica la causa y propone la acción)',
        'Los atrasos y las reprogramaciones no tienen causa registrada, así que se repiten.',
        'Las fechas atrasadas se ven en rojo en GIO y se revisan en la reunión semanal, sin clasificar la causa.',
        'Un agente clasifica cada atraso por causa (velocidad, OEE, falta de papel, cambio del cliente o saturación) y propone la '
        'acción y el ajuste de parámetros; Planificación aprueba.'),
    ('Ondupack Almendralejo', 'INT-03'): (
        'Agente de urgencias y cambios', 'Converting',
        '¿Alerta o cambio de ruta o fecha? → la ejecuta un agente (calcula el impacto) · Replanificar ruta o fecha → la ejecuta '
        'un agente (solo una parte: propone)',
        'Los cambios del cliente se aceptan por email hasta que el pedido se fabrica y las urgencias se resuelven llamando, sin '
        'registrar la causa.',
        'Planificación recibe el cambio, lo encaja a mano en la secuencia de GIO y recalcula las fechas de canal y troquelado.',
        'Un agente calcula el impacto de cada cambio o urgencia (pedidos desplazados, cambios extra, saturación del intermedio) y '
        'propone la decisión según el criterio de urgencia.'),
    ('MGT', 'INT-01'): (
        'Agente de escenarios de ciclo', 'Carga de Capacidad',
        'Carga a capacidad infinita · 3 turnos a velocidad máxima → la ejecuta un agente (solo una parte: propone el ciclo y '
        'su duración)',
        'Se carga a capacidad infinita, la máquina trabaja siempre a 3 turnos y la duración de cada calidad se ajusta sobre la '
        'marcha.',
        'Joan Carles fija a mano el ciclo de las 4 calidades y su duración con la cartera, y revisa con Comercial los pedidos que '
        'no caben.',
        'Un agente simula la duración y el orden de las calidades, las paradas y el stock, y propone el ciclo con su coste; '
        'Planificación decide.'),
    ('MGT', 'INT-02'): (
        'Agente de stock intercompany', 'Plan Fabricación',
        'Nueva tarea «Revisión mensual del stock intercompany» → la ejecuta un agente',
        'Se acumula stock de Ondupack en MGT, sin límite ni revisión.',
        'Ondupack retira poco a poco lo fabricado y nadie revisa la cobertura de sus referencias en MGT.',
        'Un agente revisa cada mes la cobertura de las referencias de Ondupack en MGT, propone mínimos y máximos y avisa del stock '
        'sin salida.'),
    ('MGT', 'INT-04'): (
        'Copiloto de plazos para Comercial', 'Demanda Externa',
        'Revisar capacidad para la fecha pedida → la ejecuta un agente (solo una parte: responde la consulta de plazo)',
        'No se propone fecha al entrar el pedido: Joan Carles revisa la capacidad a mano y además hace de atención al cliente.',
        'El cliente pide una fecha y Joan Carles comprueba a mano si hay capacidad y le responde.',
        'Comercial pregunta al copiloto cuándo sale la próxima tirada de una calidad y ancho y responde al cliente sin pasar por '
        'Planificación.'),
    ('MGT', 'INT-03'): (
        'Agente de compras de materiales', 'MRP',
        'MRP de materias primas en Excel → la ejecuta un agente (solo una parte: revisa los mensajes del MRP) · Solicitud de '
        'compra y pedido → la ejecuta un agente (solo una parte: prepara la solped)',
        'El MRP de materias primas vive en un Excel, el de auxiliares se descuadra y las solicitudes de compra se preparan a '
        'mano.',
        'Marta Boyo calcula las materias primas en Excel con la planificación de Joan Carles; Joan Carles lanza el MRP de '
        'auxiliares en SAP y crea las solped y los pedidos.',
        'Un agente revisa cada día los mensajes del MRP, propone las solicitudes de compra y avisa de los consumos que se separan '
        'del plan; Marta Boyo y Joan Carles aprueban.'),
    ('MGT', 'INT-05'): (
        'Agente de urgencias', 'Programación',
        'Encajar la urgencia a mano → la ejecuta un agente (solo una parte: propone si cabe y qué mueve)',
        'Las urgencias, también las de Ondupack fuera de plazo, se encajan a mano, sin alertas ni reglas.',
        'Joan Carles revisa a mano si la urgencia cabe en lo planificado y, si hay varias comprometidas, lo negocia con Comercial.',
        'Un agente comprueba si la urgencia cabe en el ciclo y qué pedidos y combinaciones afecta, y propone la decisión según las '
        'reglas.'),
    ('MGT', 'INT-06'): (
        'Agente de coordinación de cargas con Ondupack', 'Entregas a Ondupack y Cierre',
        'Enviar el stock a Ondupack cada día → la ejecuta un agente (toda la tarea) · Ondupack decide qué se carga al día '
        'siguiente → la ejecuta un agente (solo una parte: propone la carga)',
        'La coordinación diaria de entregas con Ondupack lleva unos 16 correos al día.',
        'Cada día Joan Carles envía el stock a Ondupack por correo y Ondupack contesta qué cargar al día siguiente.',
        'Un agente lee las peticiones de Ondupack, comprueba stock y plan y propone la carga del día dentro del mínimo y máximo; '
        'Joan Carles aprueba.'),
}


# sección 3 del formulario de agente: cómo se llama cada tarea después y qué tareas nuevas añade (y quién las hace)
DESPUES = {
    ('Ondupet', 'INT-01'): ('Comprobación de Accuracy → «Desviaciones detectadas por el agente» · Corrección Previsiones Futuros '
                            'Meses → «Aprobar la corrección propuesta»',
                            'Revisar las familias que el agente marca como dudosas — persona (Planificación)'),
    ('Ondupet', 'INT-02'): ('Decisión de turnos y balanceo → «Aprobar el plan de turnos propuesto»',
                            'Revisar los escenarios que el agente marca como de riesgo — persona (Planificación)'),
    ('Ondupet', 'INT-03'): ('Cartera Pedidos → «Revisar fechas y notas de carga propuestas»',
                            'Confirmar la nueva fecha al cliente — persona (Comercial, por O2C)'),
    ('Ondupet', 'INT-04'): ('¿Cambios de última hora? → «Impacto de la urgencia calculado» · Recuenciar urgencias → «Aprobar la '
                            'secuencia propuesta»',
                            'Decidir los conflictos que el agente escala — persona (dirección de planta)'),
    ('Ondupack Almendralejo', 'ACC-07'): ('Necesidad de papel → «Faltas por orden con acción propuesta»',
                                          'Aprobar la acción ante cada falta — persona (Planificación y Compras)'),
    ('Ondupack Almendralejo', 'INT-01'): ('Replanificar ruta, fecha o turnos → «Aprobar el plan de capacidad propuesto»',
                                          'Revisar los escenarios de riesgo de la 924 — persona (Planificación con Producción)'),
    ('Ondupack Almendralejo', 'INT-02'): ('Seguimiento de fechas atrasadas → «Atrasos con causa y acción propuesta»',
                                          'Aprobar la acción y el ajuste de parámetros — persona (Planificación)'),
    ('Ondupack Almendralejo', 'INT-03'): ('Replanificar ruta o fecha → «Aprobar el cambio con su impacto»',
                                          'Decidir los conflictos que el agente escala — persona (Planificación con Comercial)'),
    ('MGT', 'INT-01'): ('Carga a capacidad infinita → «Ciclo propuesto con su coste»',
                        'Aprobar el ciclo — persona (Joan Carles con Comercial)'),
    ('MGT', 'INT-02'): ('Revisión mensual del stock intercompany → la ejecuta el agente',
                        'Aprobar mínimos y máximos — persona (Joan Carles con Angélica, Ondupack)'),
    ('MGT', 'INT-04'): ('Revisar capacidad para la fecha pedida → «Consulta de plazo resuelta por el copiloto»',
                        'Atender las consultas que el copiloto no resuelve — persona (Joan Carles)'),
    ('MGT', 'INT-03'): ('MRP de materias primas en Excel → «Mensajes del MRP revisados por el agente» · Solicitud de compra y '
                        'pedido → «Aprobar la solped propuesta»',
                        'Revisar los consumos fuera de plan que marca el agente — persona (Marta Boyo)'),
    ('MGT', 'INT-05'): ('Encajar la urgencia a mano → «Aprobar la propuesta del agente»',
                        'Decidir con Comercial los conflictos que el agente escala — persona (Joan Carles)'),
    ('MGT', 'INT-06'): ('Enviar el stock a Ondupack cada día → desaparece · Ondupack decide qué se carga → «Aprobar la carga '
                        'propuesta»',
                        'Resolver las peticiones de Ondupack que no encajan en el mínimo y máximo — persona (Joan Carles)'),
}
# lo que ya está escrito en la aplicación y se respeta
RESULTADO_APP = {('Ondupet', 'INT-01'): 'Disminución del % de roturas y sobrestock que tenemos actualmente'}


def _fechas(empresa, base):
    return base if isinstance(base, tuple) else FECHAS[empresa][base][:2]


def hojas_fichas(wb):
    wp = wb.create_sheet('Fichas de palanca')
    wn = wb.create_sheet('Interv. negocio y software')
    wa = wb.create_sheet('Interv. agentes')

    # --- intervenciones de negocio y software ---
    titulo(wn, 'Intervenciones de negocio y software · campos del formulario',
           'Cómo va: «pendiente» en todas. Responsable: «el owner» deja el campo vacío (lo lleva el owner de la palanca). La última '
           'columna propone marcar «Además cambia tareas del flujo» para que la intervención cuente en el to-be y en la capacidad.')
    cab = ['Empresa', 'Código', 'Nombre', 'De qué es', 'Palanca', 'Qué hay que hacer', 'Responsable', 'Empieza', 'Termina',
           'Marcar «cambia tareas del flujo» (caja › tarea → qué le pasa)', 'Nota']
    cabecera(wn, 4, cab, [20, 9, 34, 10, 9, 62, 30, 11, 11, 52, 30])
    r = 5
    for emp in EMPRESAS:
        for cod, nom, tipo, pal, _, _, desc in INTERVENCIONES[emp]:
            if (emp, cod) not in NEG:
                continue
            base, resp, flujo, nota = NEG[(emp, cod)]
            ini, fin = _fechas(emp, base)
            vals = [emp, cod, nom, tipo, pal, descripcion(emp, desc), resp, ini, fin, flujo or None, nota or None]
            for j, v in enumerate(vals, 1):
                celda(wn, r, j, v, centrada=j in (2, 4, 5, 8, 9), fmt='dd/mm/yyyy' if j in (8, 9) else None)
            r += 1
    n_ult = r - 1
    wn.freeze_panes = 'D5'
    wn.auto_filter.ref = f'A4:K{n_ult}'

    # --- agentes ---
    titulo(wa, 'Intervenciones de agentes · campos del formulario',
           'Qué le pasa a la caja: «reduce» (−30 % de dedicación por defecto; se ajusta en el lienzo). Confianza: EST. Responsable '
           'de la intervención: vacío, lo lleva el owner. Problema, cómo funciona hoy y resultado de INT-01 en Ondupet son los que ya escribiste.')
    cab = ['Empresa', 'Código', 'Nombre', 'Palanca', 'Caja', 'Tareas del as-is y qué les pasa', 'Problema', 'Cómo funciona hoy',
           'Qué queremos cambiar (para el comité)', 'Después: cómo se llama la tarea', 'Tareas nuevas que añade (quién la hace)',
           'Cobertura objetivo %', 'Excepción humana objetivo %', 'Resultado esperado', 'Criterio de éxito', 'Desde', 'Hasta (go-live)']
    cabecera(wa, 4, cab, [20, 9, 30, 9, 20, 58, 46, 42, 52, 46, 42, 11, 11, 40, 44, 11, 11])
    r = 5
    for emp in EMPRESAS:
        for cod, nom, tipo, pal, _, _, _ in INTERVENCIONES[emp]:
            if (emp, cod) not in AGENTES:
                continue
            prop, caja, tareas, problema, hoy, comite = AGENTES[(emp, cod)]
            ini, fin, it = FECHAS[emp][prop]
            despues, nuevas = DESPUES[(emp, cod)]
            vals = [emp, cod, nom, pal, caja, tareas, problema, hoy, comite, despues, nuevas, it['cob'], it['exc'],
                    RESULTADO_APP.get((emp, cod), it['res']), it['crit'], ini, fin]
            for j, v in enumerate(vals, 1):
                celda(wa, r, j, v, centrada=j in (2, 4, 12, 13, 16, 17),
                      fmt='0%' if j in (12, 13) else 'dd/mm/yyyy' if j in (16, 17) else None)
            r += 1
    a_ult = r - 1
    wa.freeze_panes = 'D5'
    wa.auto_filter.ref = f'A4:Q{a_ult}'

    # --- palancas ---
    titulo(wp, 'Fichas de palanca por empresa · pestaña «Definición»',
           'El compromiso de valor (€ P&L, working capital) se deja para la cuantificación: confianza «por cuantificar» y estado '
           '«identificada». La ventana se calcula con las fechas de sus intervenciones.')
    cab = ['Empresa', 'Palanca', 'Nombre', 'Qué ataca: líneas (tipo · hallazgo)', 'Horas de personas (caja · quién · horas/año hoy)',
           'Responsable de la palanca', 'Equipo de apoyo', 'Ventana: desde', 'Ventana: hasta', 'Contador', 'Proyecto',
           'Qué cambiar en la aplicación']
    cabecera(wp, 4, cab, [20, 9, 30, 44, 48, 28, 40, 12, 12, 44, 30, 52])
    nombres = {'P2R-01': 'Gestión de la demanda', 'P2R-02': 'S&OP y plan de capacidad', 'P2R-03': 'Política de inventario por segmento',
               'P2R-04': 'Promesa de fecha fiable', 'P2R-05': 'Planificación de materiales por excepción',
               'P2R-06': 'Programación estable y optimizada', 'P2R-07': 'Liberación controlada',
               'P2R-08': 'Control del servicio y mejora continua'}
    qn = "'Interv. negocio y software'"
    qa = "'Interv. agentes'"
    rn = lambda col: f'{qn}!${col}$5:${col}${n_ult}'
    ra = lambda col: f'{qa}!${col}$5:${col}${a_ult}'
    r = 5
    for emp in EMPRESAS:
        for pid in PALANCAS[emp]:
            lineas = []
            for hid, _, famnat, _, _, pals in HALLAZGOS[emp]:
                if pid in pals and (emp, pid, hid) not in QUITAR:
                    nat = FAMILIA_CORREGIDA.get((emp, hid), famnat.split(' · ')[1])
                    lineas.append(f'{LINEA[nat]} · {hid}')
            owner, equipo, horas, contador, cambiar = FICHA_PALANCA[(emp, pid)]
            cn = f'COUNTIFS({rn("A")},$A{r},{rn("E")},$B{r})'
            ca = f'COUNTIFS({ra("A")},$A{r},{ra("D")},$B{r})'
            desde = (f'=MIN(IF({cn}>0,_xlfn.MINIFS({rn("H")},{rn("A")},$A{r},{rn("E")},$B{r}),2958465),'
                     f'IF({ca}>0,_xlfn.MINIFS({ra("P")},{ra("A")},$A{r},{ra("D")},$B{r}),2958465))')
            hasta = (f'=MAX(_xlfn.MAXIFS({rn("I")},{rn("A")},$A{r},{rn("E")},$B{r}),'
                     f'_xlfn.MAXIFS({ra("Q")},{ra("A")},$A{r},{ra("D")},$B{r}))')
            vals = [emp, pid, nombres[pid], '\n'.join(lineas), horas, owner, equipo, desde, hasta, contador, PROYECTO[emp],
                    cambiar or None]
            for j, v in enumerate(vals, 1):
                celda(wp, r, j, v, centrada=j in (2, 8, 9), fmt='dd/mm/yyyy' if j in (8, 9) else None)
            r += 1
    p_ult = r - 1
    wp.conditional_formatting.add(f'L5:L{p_ult}', FormulaRule(formula=['LEN($L5)>0'], fill=fill('FCE4E4')))
    wp.freeze_panes = 'D5'
    wp.auto_filter.ref = f'A4:L{p_ult}'
    for ws in (wp, wn, wa):
        ws.sheet_view.showGridLines = False
    # orden: palancas antes que sus intervenciones
    return wp, wn, wa
