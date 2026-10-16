"""Genera el maestro genérico de hallazgos y palancas de Plan to Release (P2R).

Dos niveles:
- Nivel alto (hojas Hallazgos, Palancas y Relación): lo que se da de alta en el maestro.
- Detalle (hojas «Detalle hallazgos» y «Detalle palancas»): la versión anterior, ya depurada,
  que se usará para bajar a intervenciones (acciones, automatizaciones y agentes).

Uso: python generar_maestro_p2r.py [salida.xlsx]
"""
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

F = 'Arial'
NAVY, TEAL, GREY, BORDER = '00263A', '244C5A', '5B6B78', 'D7DFE5'
OPOR, WARN = 'FFF1D6', 'FCE4E4'

FAM_I, FAM_O = 'Ineficiencia (coste)', 'Oportunidad (ingreso expuesto)'
ESF, FUGA, CIRC, TRANS, ING = ('Esfuerzo (ya contado dentro de personas)', 'Fuga de P&L', 'Circulante × CMPC',
                               'Transformación (horas de línea, mermas)', 'Ingreso expuesto')
FASES = ['Demanda y previsión', 'S&OP y Plan Maestro', 'Necesidades y Aprovisionamiento', 'Programación',
         'Liberación', 'Seguimiento y Control']

# ---------------------------------------------------------------- nivel alto
# (id, fase, hallazgo, descripción, familia, naturaleza, métrica, palancas, nota)
HALLAZGOS = [
    ('H01', 1, 'Previsión inestable',
     'La previsión no se cierra ni se consensúa, cambia dentro del mes en curso y su horizonte no cubre las decisiones que alimenta.',
     FAM_I, CIRC, '€ de previsión mensual cambiada tras el cierre o sin consensuar', ['PL01'], ''),
    ('H02', 1, 'Previsión sin medición ni aprendizaje',
     'No se miden precisión ni sesgo, la demanda perdida no se registra y las desviaciones no llevan acción.',
     FAM_I, CIRC, 'Familias fuera de umbral sin acción / mes', ['PL01'], ''),
    ('H03', 2, 'Demanda y capacidad sin equilibrar',
     'Sin S&OP, la capacidad (turnos, reservas, intercompany) se decide por intuición y a mano, y el desequilibrio se arregla con urgencias.',
     FAM_I, TRANS, 'Horas de línea y turnos no ajustados a la carga / mes', ['PL02'],
     'Si pesan más las horas extra que las horas de línea, usa Fuga de P&L.'),
    ('H04', 2, 'Stock sin política',
     'Sin segmentación ni estrategia por material, sobra stock de unos y falta de otros, en producto terminado y en materias primas.',
     FAM_I, CIRC, '€ de stock fuera de política (exceso y obsoleto)', ['PL03'],
     'La política la decide Planificación; el dato se guarda en R2M.'),
    ('H05', 2, 'Fecha al cliente no fiable',
     'La fecha se promete sin comprobar disponibilidad, no se renegocia cuando no se llega y hay roturas en lo que debería estar en stock.',
     FAM_O, ING, '€ de pedidos / mes entregados tarde o no servidos', ['PL04', 'PL03'],
     'Comparte ingreso con H11: cuantifica solo uno por celda.'),
    ('H06', 3, 'Faltas tardías y gestionadas a mano',
     'El MRP corre con poca frecuencia, las faltas se revisan todas a mano y acaban en compras y transportes urgentes.',
     FAM_I, FUGA, '€ de sobrecoste urgente por faltas / mes', ['PL05'], ''),
    ('H07', 3, 'Necesidades poco fiables',
     'Parámetros desactualizados, producto terminado y semielaborado sin sincronizar y agrupación de lotes a mano.',
     FAM_I, CIRC, '€ de stock de semielaborado y materia prima por encima de la necesidad', ['PL05'], ''),
    ('H08', 4, 'Secuencia sin optimizar',
     'Se secuencia a mano, sin capacidad finita y con la ruta por defecto.',
     FAM_I, TRANS, 'Horas de cambio y de sobrecarga / semana', ['PL06'], ''),
    ('H09', 4, 'Programa inestable',
     'Sin zona congelada ni reglas de prioridad, las urgencias obligan a resecuenciar.',
     FAM_I, TRANS, 'Resecuenciaciones / semana', ['PL06'], ''),
    ('H10', 5, 'Liberación sin control',
     'Se libera sin confirmar material, utillaje y aprovisionamiento a línea, y sin ventana de liberación.',
     FAM_I, TRANS, 'Arranques fallidos / mes', ['PL07'], ''),
    ('H11', 6, 'Servicio sin control',
     'No hay KPIs formales de servicio ni de adherencia al programa, y los atrasos se ven cuando ya han ocurrido.',
     FAM_O, ING, '€ de pedidos atrasados sin aviso / mes', ['PL08'],
     'Comparte ingreso con H05: cuantifica solo uno por celda.'),
    ('H12', 6, 'Ciclo sin aprendizaje',
     'Las desviaciones no tienen análisis de causa ni vuelven a la previsión ni a los parámetros, y el seguimiento se hace a mano.',
     FAM_I, TRANS, 'Desviaciones repetidas sin causa / mes', ['PL08'], ''),
]

# (id, palanca, qué cambia, KPI)
PALANCAS = [
    ('PL01', 'Gestión de la demanda',
     'La previsión se cierra en fecha, se consensúa cada mes entre Comercial y Planificación, cubre el horizonte de las decisiones y se mide y corrige.',
     'Precisión y sesgo por familia · cambios tras el cierre'),
    ('PL02', 'S&OP y plan de capacidad',
     'Demanda y capacidad se equilibran cada mes, con datos y con dirección: turnos, reservas de capacidad e intercompany se deciden con escenarios, no con urgencias.',
     'Ocupación real frente a planificada · horas extra no planificadas'),
    ('PL03', 'Política de inventario por segmento',
     'Cada material, de producto terminado y de materia prima, tiene segmento, estrategia y stock objetivo, revisados periódicamente.',
     '€ de stock fuera de política · roturas'),
    ('PL04', 'Promesa de fecha fiable',
     'La fecha al cliente se confirma con disponibilidad real y se renegocia en cuanto se sabe que no se llega.',
     'OTIF a fecha confirmada'),
    ('PL05', 'Planificación de materiales por excepción',
     'Las necesidades se calculan a tiempo, con parámetros vigentes y en todos los niveles, y el planificador solo atiende las excepciones críticas.',
     '€ urgentes por faltas · faltas detectadas a tiempo'),
    ('PL06', 'Programación estable y optimizada',
     'El programa respeta la capacidad, minimiza los cambios y solo se modifica dentro de reglas: zona congelada y prioridades.',
     'Horas de cambio · adherencia al programa'),
    ('PL07', 'Liberación controlada',
     'Una orden solo se libera con material, utillaje y aprovisionamiento a línea confirmados, y dentro de su ventana.',
     'Arranques fallidos'),
    ('PL08', 'Control del servicio y mejora continua',
     'Servicio y adherencia se miden solos, los riesgos se avisan antes y cada desviación tiene causa y una acción que vuelve a la planificación.',
     'OTIF · desviaciones con causa'),
]

# ---------------------------------------------------------------- detalle (versión anterior depurada)
# (id, fase, hallazgo, familia, naturaleza, métrica, nivel alto, nota)
DETALLE_H = [
    ('H02', 1, 'La previsión no tiene día de cierre y se cambia dentro del mes en curso', FAM_I, CIRC,
     'Cambios de previsión dentro del mes en curso / mes', 'H01', ''),
    ('H03', 1, 'No hay consenso de demanda entre Comercial y Planificación, o hay reunión pero no decide', FAM_I, CIRC,
     '€ de previsión mensual sin consensuar', 'H01', ''),
    ('H05', 1, 'La demanda perdida (agotados, pedidos rechazados) no se registra y el histórico sale sesgado', FAM_I, CIRC,
     'Agotados o rechazos / mes sin registrar', 'H02', 'Sin este dato, la previsión estadística (PL01) aprende el sesgo.'),
    ('H06', 1, 'No se miden precisión ni sesgo, o se miden sin umbral', FAM_I, CIRC, 'Familias sin KPI de precisión', 'H02', ''),
    ('H07', 1, 'Las desviaciones de previsión no llevan análisis de causa ni acción', FAM_I, CIRC,
     'Desviaciones fuera de umbral sin acción / mes', 'H02', ''),
    ('H08', 1, 'El horizonte de previsión es más corto que el plazo de las decisiones que alimenta (turnos, compras de plazo largo, intercompany)',
     FAM_I, TRANS, '€ de compras o turnos decididos sin previsión', 'H01', ''),
    ('H09', 2, 'No hay S&OP: el desequilibrio entre demanda y capacidad no se decide con dirección y se arregla abajo, con urgencias',
     FAM_I, FUGA, 'Horas extra y turnos no planificados / mes', 'H03',
     'Puede solapar con H10 y H29 (mismas horas extra): cuantificar una vez.'),
    ('H10', 2, 'La capacidad no se planifica con datos: turnos y balanceo se deciden por intuición, sin cruzar carga con capacidad',
     FAM_I, TRANS, 'Horas de línea / mes', 'H03', ''),
    ('H12', 2, 'No hay política de stock por segmento: ABC/XYZ, MTS/MTO y stock de seguridad sin definir o sin revisar', FAM_I, CIRC,
     'Materiales sin estrategia vigente', 'H04', 'La política la decide Planificación (P2R); el dato se guarda en R2M.'),
    ('H13', 2, 'Stock de producto terminado excesivo u obsoleto por fabricar contra una previsión inflada', FAM_I, CIRC,
     '€ de stock sin movimiento > 6 meses', 'H04', 'Si el stock se deprecia o se destruye, esa pérdida es Fuga de P&L.'),
    ('H14', 2, 'Roturas en materiales contra stock: hay pedidos que no se pueden servir', FAM_O, ING,
     '€ de pedidos no servidos / mes', 'H05', ''),
    ('H15', 2, 'La fecha al cliente se promete sin comprobar stock ni capacidad', FAM_O, ING,
     '€ de pedidos / mes confirmados con fecha no viable', 'H05', ''),
    ('H16', 2, 'Cuando se sabe que no se llega, la fecha no se renegocia con el cliente', FAM_O, ING,
     '€ de pedidos / mes entregados tarde sin aviso', 'H05', ''),
    ('H18', 2, 'Cálculos repetitivos a mano: reservas de capacidad (temporales), coberturas, cargas', FAM_I, ESF,
     'Horas / mes', 'H03', 'Es un cálculo fijo: mejor una regla automática que un agente.'),
    ('H19', 2, 'Intercompany (opcional): la planta que suministra se entera por un pedido urgente, sin plan conjunto ni mín/máx',
     FAM_I, TRANS, 'Pedidos intercompany urgentes / mes', 'H03', 'Solo para las divisiones que se suministran entre sí.'),
    ('H21', 3, 'Parámetros de planificación desactualizados: lote mínimo y múltiplo, plazos, stock de seguridad', FAM_I, CIRC,
     '% de materiales sin revisar en 12 meses', 'H07', ''),
    ('H22', 3, 'El MRP corre con poca frecuencia y las faltas se detectan tarde', FAM_I, TRANS,
     'Faltas / mes detectadas con menos margen que el plazo del proveedor', 'H06', ''),
    ('H23', 3, 'Las faltas se revisan todas a mano, sin criticidad ni gestión por excepción', FAM_I, ESF,
     'Líneas de falta revisadas / semana', 'H06', ''),
    ('H24', 3, 'Compras y transportes urgentes por faltas no anticipadas', FAM_I, FUGA, '€ de sobrecoste urgente / mes', 'H06', ''),
    ('H25', 3, 'Producto terminado y semielaborado se planifican por separado, sin sincronizar', FAM_I, CIRC,
     '€ de stock de semielaborado', 'H07', ''),
    ('H26', 3, 'La agrupación de necesidades en lotes o combinaciones se hace a mano', FAM_I, TRANS,
     '% de merma de combinación', 'H07', ''),
    ('H27', 3, 'Materias primas y auxiliares sin política de stock: se compra a ojo o por el mínimo del proveedor', FAM_I, CIRC,
     '€ de materia prima sin consumo > 6 meses', 'H04', ''),
    ('H28', 4, 'La secuencia se hace a mano, sin optimizar los cambios (molde, formato, altura, color)', FAM_I, TRANS,
     'Horas de cambio / semana', 'H08', ''),
    ('H29', 4, 'Se programa sin capacidad finita: se carga más de lo que cabe', FAM_I, TRANS,
     'Horas cargadas por encima de la capacidad / semana', 'H08', ''),
    ('H30', 4, 'El programa no tiene zona congelada y los cambios de última hora obligan a resecuenciar', FAM_I, TRANS,
     'Resecuenciaciones / semana', 'H09', ''),
    ('H31', 4, 'Urgencias y prioridades sin reglas: se decide por llamada o se escala a dirección', FAM_I, ESF,
     'Escalados / semana', 'H09', ''),
    ('H32', 4, 'Se usa siempre la ruta por defecto, sin reequilibrar con la carga', FAM_I, TRANS,
     'Horas en ruta no óptima / mes', 'H08', ''),
    ('H34', 5, 'Se libera sin comprobar material en línea, utillaje, ficha y personal', FAM_I, TRANS, 'Arranques fallidos / mes', 'H10', ''),
    ('H35', 5, 'No hay ventana de liberación: se libera todo de golpe o en el último momento', FAM_I, CIRC,
     'Días entre liberación e inicio', 'H10', ''),
    ('H36', 5, 'El aprovisionamiento a línea no va sincronizado con la secuencia', FAM_I, TRANS,
     'Minutos de espera de material por orden', 'H10', ''),
    ('H38', 6, 'No hay KPIs formales de servicio (OTIF a fecha confirmada) ni de adherencia al programa', FAM_O, ING,
     'Pedidos / mes sin medir', 'H11', 'Habilitador: no suma € por sí solo. Piso «sin declarar».'),
    ('H39', 6, 'Los atrasos se detectan cuando ya han ocurrido, sin aviso de pedidos en riesgo', FAM_O, ING,
     '€ de pedidos atrasados / mes', 'H11', ''),
    ('H40', 6, 'Las desviaciones del plan no tienen análisis de causa y lo aprendido no vuelve a la previsión ni a los parámetros',
     FAM_I, TRANS, 'Desviaciones sin causa / mes', 'H12', ''),
    ('H41', 6, 'Los informes de seguimiento se hacen a mano', FAM_I, ESF, 'Horas / semana', 'H12', ''),
]

# (id, palanca, tipo principal, tipo secundario, evoluciona a IA, hallazgos detalle, KPI, intervención tipo, nivel alto, nota)
DETALLE_P = [
    ('PL01', 'Previsión estadística por familia y material, a 12–18 meses', 'Auto', '', 'Sí', ['H08'], 'Precisión (1–WMAPE)',
     'Modelo por familia y material con histórico limpio (incluida la demanda perdida) y estacionalidad; Comercial solo revisa las excepciones antes del cierre.',
     'PL01', ''),
    ('PL02', 'Calendario de demanda con día de cierre y mes en curso congelado', 'Política', '', 'No', ['H02'], 'Cambios tras el cierre',
     'Publicar el calendario anual con el día de cierre y bloquear en el sistema la previsión del mes en curso; los cambios posteriores, solo con aprobación y motivo.',
     'PL01', ''),
    ('PL03', 'Revisión mensual de demanda con consenso', 'Política', '', 'No', ['H03'], 'Consenso en fecha',
     'Reunión mensual Comercial–Planificación con orden del día fijo, acta y decisiones.', 'PL01', ''),
    ('PL04', 'Registro de la demanda perdida con su motivo', 'Dato', '', 'No', ['H05'], '% de agotados registrados',
     'Definir qué es un agotado o un rechazo y registrarlo a diario en el sistema, con su motivo.', 'PL01', ''),
    ('PL05', 'KPI de precisión y sesgo por familia, con umbral y plan de acción', 'Auto', '', 'No', ['H06', 'H07'],
     'Familias fuera de umbral con acción',
     'Cálculo automático mensual por familia, umbral por segmento y plan de acción para cada familia fuera de umbral.', 'PL01', ''),
    ('PL06', 'S&OP mensual con escenarios valorados en € y reunión ejecutiva', 'Política', '', 'No', ['H09', 'H19'],
     'Horas extra no planificadas · cumplimiento del plan',
     'Ciclo mensual: revisión de demanda, revisión de suministro y reunión ejecutiva; cada escenario (turnos, stock, subcontratación, intercompany) valorado en €.',
     'PL02', ''),
    ('PL07', 'Plan de capacidad: carga frente a capacidad por línea, con regla de turnos', 'Auto', '', 'Sí', ['H10'],
     'Ocupación real frente a la planificada',
     'Cruce automático de previsión cerrada × rutas × tiempos de ciclo contra el calendario; umbrales de ocupación para abrir o cerrar turnos.',
     'PL02', ''),
    ('PL08', 'Zona congelada en el programa', 'Política', '', 'No', ['H30'], 'Cambios en zona congelada',
     'Definir el horizonte congelado del programa; las excepciones se aprueban y quedan registradas.', 'PL06', ''),
    ('PL09', 'Segmentación ABC/XYZ y política de stock por segmento, revisada cada trimestre', 'Política', 'Auto', 'No',
     ['H12', 'H13', 'H14', 'H27'], 'Cobertura · roturas · stock obsoleto',
     'Criterios ABC por valor y XYZ por variabilidad, matriz que asigna MTS/MTO y stock de seguridad, y recálculo trimestral automático.',
     'PL03', ''),
    ('PL10', 'Promesa de fecha con disponibilidad real: stock y capacidad libre (ATP/CTP)', 'Auto', '', 'No', ['H15'],
     '% de pedidos confirmados con fecha viable',
     'Consulta de disponibilidad (stock y capacidad libre) al confirmar el pedido: la fecha la propone el sistema, no se teclea.',
     'PL04', ''),
    ('PL11', 'Política de fecha: dos fechas en el pedido, plazo de confirmación y aviso al cliente', 'Política', '', 'No', ['H16'],
     'OTIF a fecha confirmada',
     'Fecha solicitada y fecha confirmada en cada pedido; plazo fijo de confirmación y aviso al cliente si no se llega.', 'PL04', ''),
    ('PL12', 'Automatizar cálculos e informes repetitivos', 'Auto', '', 'No', ['H18', 'H41'], 'Horas liberadas',
     'Pasar a reglas del sistema los cálculos e informes que hoy se hacen en Excel cada día, semana o mes.', 'PL02',
     'Los informes de seguimiento (H41) van con la palanca de nivel alto {PL08}.'),
    ('PL13', 'Plan intercompany con mín/máx y pedidos programados', 'Política', '', 'No', ['H19'], 'Pedidos intercompany urgentes',
     'Mín/máx por material entre plantas, pedidos programados y revisión semanal conjunta.', 'PL02', ''),
    ('PL14', 'Necesidades por excepción: alertas por criticidad', 'Auto', '', 'Sí', ['H23'], 'Líneas de falta revisadas',
     'El sistema marca solo las excepciones, ordenadas por criticidad.', 'PL05', ''),
    ('PL15', 'Revisión periódica de los parámetros de planificación contra la realidad', 'Dato', '', 'No', ['H21'],
     '% de parámetros vigentes',
     'Revisión semestral de lotes, plazos y stock de seguridad contra los datos reales, con un responsable por parámetro.', 'PL05', ''),
    ('PL16', 'MRP diario con alertas tempranas de faltas', 'Auto', '', 'No', ['H22', 'H24'], 'Faltas detectadas a tiempo · € urgentes',
     'MRP diario (nocturno) con alerta de faltas priorizada por fecha de necesidad y por pedido de cliente.', 'PL05', ''),
    ('PL17', 'Planificación multinivel sincronizada: producto terminado, semielaborado y materias primas', 'Sistema', '', 'No', ['H25'],
     'Stock de semielaborado',
     'Planificar los tres niveles en el mismo cálculo, con trazabilidad de cada necesidad hasta su origen.', 'PL05', ''),
    ('PL18', 'Optimizador de agrupación y combinación', 'Algoritmo', '', 'No', ['H26'], '% de merma',
     'Algoritmo de agrupación por reglas de lote y combinación que minimiza la merma y los cambios.', 'PL05', ''),
    ('PL19', 'Secuenciador con capacidad finita, matriz de cambios y rutas alternativas', 'Algoritmo', '', 'Sí',
     ['H28', 'H29', 'H32'], 'Horas de cambio · adherencia',
     'Matriz de tiempos de cambio por línea y capacidad finita; propone la secuencia y las rutas alternativas, y el programador la aprueba.',
     'PL06', ''),
    ('PL20', 'Reglas de prioridad y de escalado', 'Política', '', 'No', ['H31'], 'Escalados / semana',
     'Criterios escritos (cliente, fecha, margen, penalización) y un árbitro definido para los conflictos.', 'PL06', ''),
    ('PL22', 'Semáforo de liberación 24–72 h antes, con ventana de liberación', 'Auto', '', 'No', ['H34', 'H35', 'H36'],
     'Arranques fallidos',
     'Comprobación automática de material en línea, utillaje y personal; la orden se libera en verde y dentro de su ventana.',
     'PL07', ''),
    ('PL23', 'Cuadro de mando de servicio y adherencia, con alertas de pedidos en riesgo', 'Auto', '', 'No', ['H38', 'H39'],
     'OTIF · avisos anticipados',
     'OTIF a fecha confirmada y adherencia al programa calculados solos, con alerta de los pedidos en riesgo antes de que se retrasen.',
     'PL08', ''),
    ('PL24', 'Análisis de causa raíz de las desviaciones, que vuelve a la previsión y a los parámetros', 'Política', '', 'Sí',
     ['H07', 'H40'], 'Desviaciones con causa',
     'Clasificar cada desviación por causa (lista cerrada), asignar una acción con responsable y devolver lo aprendido a la previsión y a los parámetros.',
     'PL08', ''),
    ('PL25', 'Copiloto de consulta: «¿cuándo sirvo X?», «¿por qué va tarde X?»', 'Consulta', '', 'No', ['H15', 'H39'],
     'Consultas resueltas sin llamada',
     'Asistente que responde sobre plazos y estado de pedidos con los datos del sistema, sin modificar nada.', 'PL04',
     '«¿Por qué va tarde X?» también sirve a la palanca de nivel alto {PL08}.'),
]

# (id, proceso destino, hallazgo, familia, naturaleza, métrica, palanca en ese proceso)
OTROS = [
    ('X01', 'R2M', 'El alta de material no pasa por Planificación (estrategia, ruta, parámetros)', FAM_I, CIRC,
     'Altas / mes sin validar por Planificación', 'Validación de Planificación en el alta de material'),
    ('X02', 'R2M', 'Duplicados en el maestro, sin criterio de qué es un material', FAM_I, ESF,
     'Materiales duplicados', 'Criterio de material y auditoría de duplicados'),
    ('X03', 'R2M', 'El calendario de turnos y paradas, las velocidades y los tiempos de cambio no coinciden con la realidad',
     FAM_I, TRANS, '% de recursos con datos sin revisar', 'Maestro de recursos de fabricación mantenido'),
    ('X04', 'P2P', 'Plazos de proveedor desactualizados y confirmaciones sin registrar', FAM_I, CIRC,
     '% de líneas con fecha confirmada distinta de la prevista', 'Seguimiento de confirmaciones de proveedor'),
    ('X05', 'O2C', 'Cartera con errores (cantidades mayores que el pedido, fechas imposibles) y pedidos servidos sin saldar',
     FAM_I, CIRC, 'Líneas de cartera erróneas o servidas sin saldar', 'Validación de cargas y saldado por regla'),
]

# (n.º, hallazgo creado en Ondupet, id nivel alto, detalle, nota)
ONDUPET_H = [
    (1, 'Sin cierre de previsiones: se modifica a mitad de mes', 'H01', 'H02', ''),
    (2, 'No se hace carga de capacidad (se decide por sensaciones)', 'H03', 'H10', '2 y 4 caen en el mismo hallazgo de nivel alto.'),
    (3, 'No se calculan ABC ni estrategias', 'H04', 'H12', ''),
    (4, 'Temporales calculados a mano cada mes', 'H03', 'H18', '2 y 4 caen en el mismo hallazgo de nivel alto.'),
    (5, 'No se modifican las fechas al cliente aunque se sepa que no se llega', 'H05', 'H16', 'Familia: Oportunidad (ingreso expuesto).'),
    (6, 'Saldado de pedidos manual; a veces no se hace', 'X05', '', 'Va a O2C × Ondupet con «Detectado en: P2R».'),
    (7, 'Corrección manual de desviaciones de demanda', 'H02', '',
     'A confirmar: su detalle (antiguo H20) lo borraste; lo más cercano es H02.'),
    (8, 'Revisión de faltas manual', 'H06', 'H23', ''),
    (9, 'Secuenciación a mano', 'H08', 'H28', ''),
]
# (n.º, palanca creada en Ondupet, id nivel alto, detalle)
ONDUPET_P = [
    ('P1', 'Previsión cerrada y congelada', 'PL01', 'PL02, PL05'),
    ('P2', 'Capacidad decidida con datos', 'PL02', 'PL07, PL12'),
    ('P3', 'Segmentación y estrategias al día', 'PL03', 'PL09'),
    ('P4', 'Fecha comprometida viable', 'PL04', 'PL11, PL10'),
    ('P5', 'Saldado por regla', 'X05', ''),
    ('P6', 'Necesidades y faltas por excepción', 'PL05', 'PL14, PL16'),
    ('P7', 'Secuencia optimizada', 'PL06', 'PL19, PL08, PL20'),
]

THIN = Side(style='thin', color=BORDER)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(horizontal='center', vertical='top', wrap_text=True)


def d(texto):
    """Prefija con «D-» los ID de la versión de detalle para no confundirlos con los de nivel alto.

    Un ID entre llaves ({PL08}) es de nivel alto: se deja tal cual y se quitan las llaves."""
    return re.sub(r'(?<!\{)\b(H\d\d|PL\d\d)\b', r'D-\1', texto).replace('{', '').replace('}', '')


def fill(c):
    return PatternFill('solid', fgColor=c)


def titulo(ws, texto, sub):
    ws['A1'] = texto
    ws['A1'].font = Font(name=F, size=14, bold=True, color=NAVY)
    ws['A2'] = sub
    ws['A2'].font = Font(name=F, size=9, italic=True, color=GREY)


def tabla(ws, fila, cabeceras, anchos, filas, centradas=()):
    """Escribe cabecera y filas desde `fila`; devuelve la última fila escrita."""
    for j, (h, w) in enumerate(zip(cabeceras, anchos), 1):
        c = ws.cell(fila, j, h)
        c.font, c.fill, c.alignment, c.border = Font(name=F, size=10, bold=True, color='FFFFFF'), fill(NAVY), CENTER, BOX
        ws.column_dimensions[c.column_letter].width = w
    for i, valores in enumerate(filas, fila + 1):
        for j, v in enumerate(valores, 1):
            c = ws.cell(i, j, v)
            c.font = Font(name=F, size=10, color=NAVY)
            c.alignment = CENTER if j in centradas else WRAP
            c.border = BOX
    ultima = fila + len(filas)
    ws.freeze_panes = ws.cell(fila + 1, 1)
    ws.auto_filter.ref = f'A{fila}:{ws.cell(fila, len(cabeceras)).column_letter}{ultima}'
    return ultima


def avisar_cero(ws, col, ult):
    ws.conditional_formatting.add(f'{col}5:{col}{ult}', FormulaRule(formula=[f'${col}5=0'], fill=fill(WARN),
                                                                  font=Font(name=F, bold=True, color='B42318')))


def construir(salida):
    wb = Workbook()
    wh = wb.active
    wh.title = 'Hallazgos'
    wp, wr, wdh, wdp, wo, wn = (wb.create_sheet(n) for n in
                                ('Palancas', 'Relación', 'Detalle hallazgos', 'Detalle palancas', 'Otros procesos', 'Ondupet'))
    n_rel = sum(len(h[7]) for h in HALLAZGOS)
    h_ult, p_ult, rel_ult = 4 + len(HALLAZGOS), 4 + len(PALANCAS), 4 + n_rel
    dh_ult, dp_ult, o_ult = 4 + len(DETALLE_H), 4 + len(DETALLE_P), 4 + len(OTROS)
    rng = lambda hoja, col, ult: f"'{hoja}'!${col}$5:${col}${ult}" if ' ' in hoja else f'{hoja}!${col}$5:${col}${ult}'

    # --- Hallazgos (nivel alto) ---
    titulo(wh, 'Maestro de hallazgos · Plan to Release (P2R) · nivel alto',
           'Genérico para cualquier empresa. El paso concreto y la cuantificación se ponen al usarlo en cada celda. '
           'El detalle, para bajar a intervenciones, está en «Detalle hallazgos».')
    filas = []
    for k, (hid, fo, nom, desc, fam, nat, met, pls, nota) in enumerate(HALLAZGOS):
        r = 5 + k
        filas.append([hid, 'P2R', fo, FASES[fo - 1], nom, desc, fam, nat, met, ', '.join(pls),
                      f'=COUNTIF({rng("Relación", "C", rel_ult)},A{r})',
                      f'=COUNTIF({rng("Detalle hallazgos", "G", dh_ult)},A{r})', nota or None])
    tabla(wh, 4, ['ID', 'Proceso', 'Orden', 'Fase', 'Hallazgo', 'Descripción', 'Familia', 'Naturaleza',
                  'Métrica (volumen expuesto)', 'Palancas', 'Nº palancas', 'Nº detalle', 'Notas'],
          [7, 9, 7, 22, 28, 60, 22, 26, 32, 13, 10, 10, 36], filas, centradas=(1, 2, 3, 10, 11, 12))
    wh.conditional_formatting.add(f'A5:M{h_ult}', FormulaRule(formula=[f'$G5="{FAM_O}"'], fill=fill(OPOR)))
    avisar_cero(wh, 'K', h_ult)

    # --- Palancas (nivel alto) ---
    titulo(wp, 'Maestro de palancas · Plan to Release (P2R) · nivel alto',
           'La palanca dice qué cambia. Cómo se hace (acciones, automatizaciones y agentes) se trabajará en las intervenciones, '
           'a partir de «Detalle palancas».')
    resuelve = {p[0]: [h[0] for h in HALLAZGOS if p[0] in h[7]] for p in PALANCAS}
    detalle = {p[0]: [x[0] for x in DETALLE_P if x[8] == p[0]] for p in PALANCAS}
    filas = []
    for k, (pid, nom, que, kpi) in enumerate(PALANCAS):
        r = 5 + k
        filas.append([pid, 'P2R', nom, que, ', '.join(resuelve[pid]), f'=COUNTIF({rng("Relación", "A", rel_ult)},A{r})',
                      kpi, ', '.join('D-' + x for x in detalle[pid]),
                      f'=COUNTIF({rng("Detalle palancas", "I", dp_ult)},A{r})'])
    tabla(wp, 4, ['ID', 'Proceso', 'Palanca', 'Qué cambia', 'Hallazgos que resuelve', 'Nº hallazgos', 'KPI',
                  'Detalle (para intervenciones)', 'Nº detalle'],
          [7, 9, 32, 66, 18, 11, 34, 30, 10], filas, centradas=(1, 2, 5, 6, 9))
    avisar_cero(wp, 'F', p_ult)

    # --- Relación (nivel alto) ---
    titulo(wr, 'Relación palanca ↔ hallazgo · nivel alto', 'Una fila por par, lista para importar.')
    filas, k = [], 0
    for p in PALANCAS:
        for hid in resuelve[p[0]]:
            r = 5 + k
            filas.append([p[0], f'=INDEX({rng("Palancas", "C", p_ult)},MATCH(A{r},{rng("Palancas", "A", p_ult)},0))', hid,
                          f'=INDEX({rng("Hallazgos", "E", h_ult)},MATCH(C{r},{rng("Hallazgos", "A", h_ult)},0))',
                          f'=INDEX({rng("Hallazgos", "D", h_ult)},MATCH(C{r},{rng("Hallazgos", "A", h_ult)},0))'])
            k += 1
    tabla(wr, 4, ['Palanca', 'Texto de la palanca', 'Hallazgo', 'Texto del hallazgo', 'Fase del hallazgo'],
          [10, 40, 10, 40, 28], filas, centradas=(1, 3))

    # --- Detalle hallazgos ---
    titulo(wdh, 'Detalle de hallazgos (para intervenciones)',
           'Versión anterior sin los que borraste. «D-» indica detalle; cada uno cuelga de un hallazgo de nivel alto.')
    pal_det = {}
    for x in DETALLE_P:
        for h in x[5]:
            pal_det.setdefault(h, []).append('D-' + x[0])
    filas = []
    for k, (hid, fo, txt, fam, nat, met, alto, nota) in enumerate(DETALLE_H):
        r = 5 + k
        filas.append(['D-' + hid, FASES[fo - 1], txt, fam, nat, met, alto,
                      f'=INDEX({rng("Hallazgos", "E", h_ult)},MATCH(G{r},{rng("Hallazgos", "A", h_ult)},0))',
                      ', '.join(pal_det.get(hid, [])), d(nota) or None])
    tabla(wdh, 4, ['ID detalle', 'Fase', 'Hallazgo (detalle)', 'Familia', 'Naturaleza', 'Métrica (volumen expuesto)',
                   'Nivel alto', 'Hallazgo de nivel alto', 'Palancas de detalle', 'Notas'],
          [10, 22, 56, 22, 26, 32, 10, 28, 18, 38], filas, centradas=(1, 7))
    wdh.conditional_formatting.add(f'A5:J{dh_ult}', FormulaRule(formula=[f'$D5="{FAM_O}"'], fill=fill(OPOR)))

    # --- Detalle palancas ---
    titulo(wdp, 'Detalle de palancas (para intervenciones)',
           'Punto de partida para traducir cada palanca en acciones, automatizaciones y agentes.')
    filas = []
    for k, (pid, nom, t1, t2, ia, hs, kpi, como, alto, nota) in enumerate(DETALLE_P):
        r = 5 + k
        filas.append(['D-' + pid, nom, t1, t2 or None, ia, ', '.join('D-' + h for h in hs), kpi, como, alto,
                      f'=INDEX({rng("Palancas", "C", p_ult)},MATCH(I{r},{rng("Palancas", "A", p_ult)},0))', d(nota) or None])
    tabla(wdp, 4, ['ID detalle', 'Palanca (detalle)', 'Tipo principal', 'Tipo secundario', 'Evoluciona a IA',
                   'Hallazgos de detalle', 'KPI', 'Intervención tipo (cómo)', 'Nivel alto', 'Palanca de nivel alto', 'Notas'],
          [10, 46, 13, 13, 11, 20, 28, 60, 10, 30, 32], filas, centradas=(1, 3, 4, 5, 9))

    # --- Otros procesos ---
    titulo(wo, 'Detectados en P2R, pero van al maestro de otro proceso',
           'Créalos en el proceso destino con «Detectado en: P2R»: el coste suma donde se arregla.')
    filas = [[x[0], x[1], 'P2R', x[2], x[3], x[4], x[5], x[6]] for x in OTROS]
    tabla(wo, 4, ['ID', 'Proceso destino', 'Detectado en', 'Hallazgo', 'Familia', 'Naturaleza',
                  'Métrica (volumen expuesto)', 'Palanca en ese proceso'],
          [7, 11, 11, 58, 22, 26, 32, 40], filas, centradas=(1, 2, 3))

    # --- Ondupet: dónde cae lo ya creado ---
    titulo(wn, 'Ondupet · dónde cae lo que ya está creado',
           'Para re-vincular los hallazgos y palancas creados directamente en la celda P2R × Ondupet.')

    def busca(col_id, hoja_alto, col_txt, ult_alto, col_otros, r):
        return (f'=IFERROR(INDEX({rng(hoja_alto, col_txt, ult_alto)},MATCH({col_id}{r},{rng(hoja_alto, "A", ult_alto)},0)),'
                f'INDEX({rng("Otros procesos", col_otros, o_ult)},MATCH({col_id}{r},{rng("Otros procesos", "A", o_ult)},0)))')

    filas = [[n, txt, alto, busca('C', 'Hallazgos', 'E', h_ult, 'D', 5 + i), d(det) or None, nota or None]
             for i, (n, txt, alto, det, nota) in enumerate(ONDUPET_H)]
    ult = tabla(wn, 4, ['Nº', 'Hallazgo creado', 'Nivel alto', 'Hallazgo de nivel alto', 'Detalle', 'Nota'],
                [8, 50, 11, 40, 12, 50], filas, centradas=(1, 3, 5))
    r0 = ult + 3
    wn.cell(r0 - 1, 1, 'Palancas').font = Font(name=F, size=12, bold=True, color=NAVY)
    filas = [[n, txt, alto, busca('C', 'Palancas', 'C', p_ult, 'H', r0 + 1 + i), d(det) or None, None]
             for i, (n, txt, alto, det) in enumerate(ONDUPET_P)]
    tabla(wn, r0, ['Nº', 'Palanca creada', 'Nivel alto', 'Palanca de nivel alto', 'Detalle', 'Nota'],
          [8, 50, 11, 40, 12, 50], filas, centradas=(1, 3))
    wn.freeze_panes = 'A5'
    wn.auto_filter.ref = None

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
    wb.save(salida)


if __name__ == '__main__':
    out = Path(sys.argv[1] if len(sys.argv) > 1 else Path(__file__).resolve().parents[1] / 'Maestro_P2R_Hallazgos_Palancas.xlsx')
    construir(out)
    print(out)
