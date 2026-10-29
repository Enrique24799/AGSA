"""Hojas por empresa del consolidado de P2R: hallazgos de la celda, palanca que los cura e intervenciones.

Los hallazgos, palancas e intervenciones son los dados de alta en la aplicación (capturas del 29/10/2026). Las
capturas solo muestran el nombre de cada intervención: la descripción sale de la propuesta con ese mismo contenido.
En MGT aún no hay intervenciones dadas de alta: se muestran las propuestas, pendientes de revisar.
"""
from openpyxl.styles import Alignment, Border, Font, Side

import generar_intervenciones_mgt as mgt
import generar_intervenciones_ondupack as opack
import generar_intervenciones_ondupet as opet
from generar_intervenciones_ondupet import BORDER, F, NAVY, cabecera, celda, fill, titulo

TINT, GRIS = 'EAF3FA', '7A8794'
DEM, SOP, NEC = 'Demanda y previsión', 'S&OP y Plan Maestro', 'Necesidades de Fabricación y MRP'
PRO, LIB, SEG = 'Programación', 'Liberación', 'Seguimiento y Control'
INEF, OPOR = 'Ineficiencia', 'Oportunidad'

# descripción de cada propuesta, por empresa (el mismo nombre puede tener textos distintos según la planta)
QUE = {emp: {it['nombre']: it['que'] for it in modulo.INTERVENCIONES}
       for emp, modulo in (('Ondupet', opet), ('Ondupack Almendralejo', opack), ('MGT', mgt))}


def descripcion(empresa, texto):
    """Si «texto» es el nombre de una propuesta de la empresa, devuelve su descripción; si no, es ya la descripción."""
    if texto in QUE[empresa]:
        return QUE[empresa][texto]
    if len(texto) < 70:
        raise KeyError(f'{empresa}: «{texto}» no es una propuesta conocida ni parece una descripción')
    return texto

# hallazgos de cada celda: (id, tipo, familia · naturaleza, anclaje, descripción, palancas que lo curan)
HALLAZGOS = {
    'Ondupet': [
        ('INF-10', 'HT-P2R-01 · Demanda sin previsión fiable', f'{INEF} · Circulante', f'Previsión y Precisión · {DEM}',
         'Previsión inestable: Comercial modifica la previsión a comienzo de mes, incluso con el material ya fabricado, sin '
         'comunicación formal ni cierre ni consenso, y nadie responde del stock que queda fabricado.', ['P2R-01']),
        ('INF-16', 'HT-P2R-01 · Demanda sin previsión fiable', f'{INEF} · Circulante', f'Previsión y Precisión · {DEM}',
         'Previsión sin medición ni aprendizaje: la reunión mensual de accuracy existe pero no genera acciones, hay '
         'desviaciones grandes sin causa ni corrección y los agotados se registran de forma irregular.', ['P2R-01']),
        ('INF-11', 'HT-P2R-03 · Demanda y capacidad sin equilibrar', f'{INEF} · Transformación', f'Plan Maestro · {SOP}',
         'Las necesidades se cargan sin comprobar la capacidad de extrusión y termoformado ni balancear entre líneas, y los '
         'temporales se calculan a mano.', ['P2R-02']),
        ('INF-13', 'HT-P2R-04 · Stock sin política clara', f'{INEF} · Circulante', f'Plan Maestro · {SOP}',
         'No se calculan ABC ni estrategias MTS/MTO aunque GIO tiene la transacción, y el stock fabricado por cambios de '
         'previsión queda sin dueño.', ['P2R-03']),
        ('OPP-02', 'HT-P2R-05 · Fecha al cliente no fiable', f'{OPOR} · —', f'Plan Maestro · {SOP}',
         'Se mantiene la fecha del cliente aunque se sabe desde el principio que no se llega, y las dos fechas del pedido no '
         'se usan.', ['P2R-03', 'P2R-04']),
        ('INF-12', 'HT-P2R-08 · Secuencia sin optimizar', f'{INEF} · Transformación', f'Programación · {PRO}',
         'La secuencia se hace totalmente a mano en GIO, sin optimizar los cambios de molde y altura y con la ruta por '
         'defecto sin reequilibrar.', ['P2R-06']),
        ('INF-14', 'HT-P2R-09 · Programa inestable', f'{INEF} · Transformación', f'Programación · {PRO}',
         'No hay zona congelada: los cambios de última hora llegan por teléfono y las prioridades se resuelven a mano; sin '
         'acuerdo, decide el director de planta.', ['P2R-06']),
        ('OPP-03', 'HT-P2R-11 · Servicio sin control', f'{OPOR} · —', f'Seguimiento y control · {SEG}',
         'No se mide el % de cumplimiento del servicio: el reporte diario de cargas atrasadas no es un KPI de OTIF y la fecha '
         'pensada para medir retrasos no se usa.', ['P2R-08']),
        ('INF-15', 'HT-P2R-12 · Ciclo sin aprendizaje', f'{INEF} · Transformación', f'Seguimiento y control · {SEG}',
         'El reporte de atrasos se envía sin análisis de causa y los ajustes se hacen pedido a pedido, sin volver a la '
         'previsión ni a los parámetros.', ['P2R-08']),
    ],
    'MGT': [
        ('INF-01', 'HT-P2R-01 · Demanda sin previsión fiable', f'{OPOR} · —', f'Demanda Intercompany · {DEM}',
         'No hay previsión de ningún tipo, ni externa ni intercompany: se planifica solo con la cartera, y no se sabe si se '
         'registran las ventas perdidas.', ['P2R-01']),
        ('INF-02', 'HT-P2R-03 · Demanda y capacidad sin equilibrar', f'{INEF} · Transformación', f'Carga de Capacidad · {SOP}',
         'Los pedidos se cargan a capacidad infinita y la máquina trabaja siempre a 3 turnos a velocidad máxima; la duración '
         'de cada calidad en el ciclo se ajusta sobre la marcha, sin propuesta.', ['P2R-02']),
        ('INF-03', 'HT-P2R-04 · Stock sin política clara', f'{INEF} · Circulante', f'Plan Fabricación · {SOP}',
         'Se acumula stock de Ondupack en MGT, que lo va retirando poco a poco; lo fabricado de más no se ajusta en el pedido '
         'y las 3 turnos generan stock aunque todo se fabrica bajo pedido.', ['P2R-03']),
        ('OPP-02', 'HT-P2R-05 · Fecha al cliente no fiable', f'{OPOR} · —', f'Plan Fabricación · {SOP}',
         'Al entrar el pedido no se propone fecha porque la planificación no está integrada con Comercial, y las fechas de '
         'Ondupack no son fiables: marcan el inicio del ciclo y después cambian.', ['P2R-03', 'P2R-04']),
        ('INF-04', 'HT-P2R-07 · Necesidades poco fiables', f'{INEF} · Circulante', f'Propuesta de Fabricación · {NEC}',
         'El MRP de materias primas se hace en un Excel; el de auxiliares, en SAP, se descuadra porque la entrada de báscula '
         'se registra después del consumo, y no se sabe cómo se planifica la bobina madre.', ['P2R-05']),
        ('INF-05', 'HT-P2R-08 · Secuencia sin optimizar', f'{INEF} · Transformación', f'Programación · {PRO}',
         'El algoritmo de combinación no tiene en cuenta lo ya fabricado ni hay fecha límite de cartera, y la agrupación por '
         'calidad y ancho es manual salvo que se use su salida.', ['P2R-06']),
        ('INF-06', 'HT-P2R-09 · Programa inestable', f'{INEF} · Transformación', f'Programación · {PRO}',
         'Las urgencias, también las de Ondupack fuera del plazo de 10 días, se encajan a mano, sin alertas ni zona '
         'congelada, y las prioridades se negocian caso a caso.', ['P2R-06']),
        ('INF-08', 'HT-P2R-10 · Liberación sin control', f'{INEF} · Transformación', f'Liberación de la Orden · {LIB}',
         'No hay faltas por orden: el planificador comprueba a mano, en varios sitios, que hay material, sin ver cuándo llega, '
         'y la orden se entrega impresa en papel.', ['P2R-07']),
        ('OPP-03', 'HT-P2R-11 · Servicio sin control', f'{OPOR} · —', f'Seguimiento del servicio · {SEG}',
         'No hay pilar de servicio ni KPI: las incidencias se resuelven por llamada y la coordinación de entregas con Ondupack '
         'lleva unos 16 correos al día.', ['P2R-08']),
        ('INF-07', 'HT-P2R-12 · Ciclo sin aprendizaje', f'{INEF} · Transformación', f'Seguimiento del servicio · {SEG}',
         'No se analizan desviaciones ni atrasos y las incidencias no dejan causa registrada.', ['P2R-08']),
    ],
    'Ondupack Almendralejo': [
        ('INF-01', 'HT-P2R-01 · Demanda sin previsión fiable', f'{INEF} · Circulante', f'Entrada de Pedidos y Fecha · {DEM}',
         'No se usa previsión aunque hay clientes con consumo repetitivo, no se sabe qué parte de la demanda es de pedido '
         'firme, de stock o intercompany, y la demanda perdida no se registra en el sistema.', ['P2R-01']),
        ('INF-06', 'HT-P2R-03 · Demanda y capacidad sin equilibrar', f'{INEF} · Transformación', f'Carga Capacidad · {SOP}',
         'La 924 no alcanza los metros planificados y arrastra pedidos lanzados. La sobresaturación se detecta tarde: por las '
         'fechas de entrega o porque se llena el almacén intermedio.', ['P2R-02']),
        ('INF-02', 'HT-P2R-06 · Faltas tardías y gestionadas a mano', f'{INEF} · Fuga de P&L', f'Papel y Auxiliares · {NEC}',
         'La falta de materia prima es su primer problema. El papel se pide por consumo pasado y los auxiliares se reponen a '
         'mano; no hay MRP ni aviso de faltas por orden, y GIO solo ve el stock de papel.', ['P2R-05']),
        ('INF-03', 'HT-P2R-07 · Necesidades poco fiables', f'{INEF} · Circulante', f'Papel y Auxiliares · {NEC}',
         'La necesidad de papel no sale del plan, sale del histórico. Las sustituciones se hacen a mano, los consumos se dan '
         'por diferencia de inventario cada mes y el maestro está incompleto (sin lote mínimo ni calendario).', ['P2R-05']),
        ('INF-04', 'HT-P2R-08 · Secuencia sin optimizar', f'{INEF} · Transformación', f'Onduladora · {PRO}',
         'La combinación de onduladora y la secuencia de converting no se optimizan; la optimización está aplazada hasta '
         'estabilizar los datos. Se repiten maestros fabricados en la misma semana.', ['P2R-06']),
        ('INF-05', 'HT-P2R-09 · Programa inestable', f'{INEF} · Transformación', f'Onduladora · {PRO}',
         'No hay zona congelada: los cambios del cliente se aceptan por email mientras el pedido no esté fabricado, las '
         'urgencias se resuelven llamando y no se registra la causa de las reprogramaciones.', ['P2R-06']),
        ('OPP-01', 'HT-P2R-11 · Servicio sin control', f'{OPOR} · —', f'Revisión Semanal · {SEG}',
         'Falta un KPI oficial de servicio: Dirección y Comercial todavía no han decidido cuál usar ni contra qué fecha '
         'medirlo.', ['P2R-08']),
    ],
}

# palancas aplicadas en cada celda, en orden
PALANCAS = {
    'Ondupet': ['P2R-01', 'P2R-02', 'P2R-03', 'P2R-04', 'P2R-06', 'P2R-08'],
    'MGT': ['P2R-01', 'P2R-02', 'P2R-03', 'P2R-04', 'P2R-05', 'P2R-06', 'P2R-07', 'P2R-08'],
    'Ondupack Almendralejo': ['P2R-01', 'P2R-02', 'P2R-05', 'P2R-06', 'P2R-08'],
}

NC = 'no cambia caja'
# intervenciones dadas de alta: (código, nombre, de qué es, palanca, caja que cambia, capacidad €, descripción o nombre de la propuesta)
INTERVENCIONES = {
    'Ondupet': [
        ('ACC-01', 'Calendario de previsión con día de cierre', 'negocio', 'P2R-01', NC, None,
         'Calendario de previsión con día de cierre'),
        ('ACC-02', 'Registro diario de agotados', 'negocio', 'P2R-01', NC, None, 'Registro diario de agotados y demanda perdida'),
        ('INT-01', 'Análisis de las desviaciones de la demanda', 'agente', 'P2R-01', 'Previsión y Precisión', 903,
         'Agente de análisis de desviaciones'),
        ('ACC-03', 'Cálculo de % de saturación', 'software', 'P2R-02', NC, None,
         'Al cerrar la previsión, el sistema calcula el % de saturación de cada línea de extrusión y termoformado (previsión × '
         'rutas × tiempos de ciclo frente al calendario) y marca las que pasan del umbral.'),
        ('ACC-04', 'Revisión mensual de la capacidad de la planta', 'negocio', 'P2R-02', NC, None,
         'Revisión mensual de capacidad con dirección de planta'),
        ('ACC-05', 'Cálculo automático de temporales', 'software', 'P2R-02', NC, None, 'Temporales automáticos'),
        ('INT-02', 'Propuestas de balanceo', 'agente', 'P2R-02', 'Plan Maestro', 6000,
         'Agente que propone el reparto de turnos y de carga entre las líneas de extrusión y termoformado, con su impacto en '
         'horas y €. Planificación lo aprueba en la revisión mensual de capacidad.'),
        ('ACC-06', 'Definir criterios ABC / estrategia', 'negocio', 'P2R-03', NC, None,
         'Criterios de segmentación y estrategia de stock'),
        ('ACC-07', 'Definir acciones para el stock sin rotación', 'negocio', 'P2R-03', NC, None,
         'Responsable del stock sin movimiento'),
        ('ACC-08', 'Cálculo automático de días de stock', 'software', 'P2R-03', NC, None,
         'El sistema calcula la cobertura en días de stock de cada referencia y avisa de las que superan el objetivo de su '
         'segmento, para decidir a tiempo qué se deja de fabricar o se liquida.'),
        ('ACC-09', 'Definir política de fecha confirmada', 'negocio', 'P2R-04', NC, None, 'Política de fecha confirmada'),
        ('INT-03', 'Propuesta de fechas de entrega viables para los pedidos y notas de carga', 'agente', 'P2R-04',
         'Plan Maestro', 6000,
         'Agente que calcula una fecha de entrega viable para cada pedido con la capacidad y los temporales, y propone las '
         'notas de carga. Comercial confirma la fecha al cliente por O2C.'),
        ('INT-04', 'Determinación de urgencias', 'agente', 'P2R-06', 'Programación', 4500, 'Agente de urgencias'),
        ('ACC-10', 'Análisis de atrasos de entregas', 'negocio', 'P2R-08', NC, None, 'Revisión mensual de causas de atraso'),
        ('ACC-11', 'Informe automático de pedidos / cargas en riesgo', 'software', 'P2R-08', NC, None,
         'Informe automático diario de los pedidos y cargas en riesgo de llegar tarde: sustituye al reporte manual de cargas '
         'atrasadas y avisa a Comercial antes de que ocurra el atraso.'),
    ],
    'Ondupack Almendralejo': [
        ('ACC-01', 'Integración de las referencias/clientes que vayan por previsión en GIO', 'negocio', 'P2R-01', NC, None,
         'Identificar con Comercial los clientes y referencias con consumo repetitivo, acordar con ellos una previsión no '
         'vinculante y cargarla en GIO para que la carga-capacidad y el CTP la tengan en cuenta.'),
        ('ACC-02', 'Registro de la demanda perdida en GIO', 'software', 'P2R-01', NC, None, 'Registro de la demanda perdida en GIO'),
        ('ACC-03', 'Revisión mensual de capacidad y cuellos de botella', 'negocio', 'P2R-02', NC, None,
         'Revisión mensual de capacidad y cuellos de botella'),
        ('INT-01', 'Agente de escenarios de capacidad', 'agente', 'P2R-02', 'Carga Capacidad', 0,
         'Agente de escenarios de capacidad'),
        ('ACC-04', 'Necesidad de papel calculada desde el plan', 'software', 'P2R-05', NC, None,
         'Necesidad de papel calculada desde el plan'),
        ('ACC-05', 'Stock de papel fiable sin esperar al inventario mensual', 'software', 'P2R-05', NC, None,
         'Stock de papel fiable sin esperar al inventario'),
        ('ACC-06', 'Previsión de papel compartida con MGT/PDA/PAPRESA', 'negocio', 'P2R-05', NC, None,
         'En cada ciclo, Ondupack envía a MGT, PDA y Papresa la previsión de papel de los próximos ciclos por calidad y '
         'ancho, además del pedido firme 10 días antes. El pedido sigue yendo por P2P.'),
        ('ACC-07', 'Faltas de materia prima', 'agente', 'P2R-05', NC, None, 'Agente de faltas de materia prima'),
        ('ACC-11', 'Algoritmo de combinación', 'software', 'P2R-06', NC, None, 'Algoritmo de combinación de onduladora'),
        ('INT-03', 'Agente de urgencias y cambios', 'agente', 'P2R-06', 'Converting', 0, 'Agente de urgencias y cambios'),
        ('ACC-08', 'KPI oficial de servicio', 'negocio', 'P2R-08', NC, None, 'KPI oficial de servicio'),
        ('ACC-09', 'Reunión semanal con KPIs y causas', 'negocio', 'P2R-08', NC, None, 'Reunión semanal con KPIs y causas'),
        ('ACC-10', 'Copiloto de estado de pedidos', 'software', 'P2R-08', NC, None, 'Copiloto de estado de pedidos'),
        ('INT-02', 'Agente de causa raíz', 'agente', 'P2R-08', 'Seguimiento del Servicio', 0, 'Agente de causa raíz'),
    ],
}
TIPO_DICE = {'Negocio': 'negocio', 'Automatización': 'software', 'Agente': 'agente', 'Consulta': 'software'}
INTERVENCIONES['MGT'] = [
    ('propuesta', it['nombre'], TIPO_DICE[it['tipo']], it['pal'], it['cajas'][0][0] if it['cajas'] else NC, None, it['que'])
    for it in mgt.INTERVENCIONES]

CABECERA = {
    'Ondupet': 'Process owner: celda sin activar. Capacidad liberada que calcula la aplicación: 17,4 k€.',
    'MGT': 'Process owner: celda sin activar. Sin intervenciones dadas de alta: las de la tabla son las propuestas, '
           'pendientes de revisar (en gris).',
    'Ondupack Almendralejo': 'Process owner: Enrique Tejada. La capacidad de los agentes sale a 0 € porque la celda aún no '
                             'tiene dedicación de personas por fase.',
}
NOTAS = {
    'Ondupet': [
        'INF-10 e INF-16 son del mismo tipo (HT-P2R-01) desde que se agruparon los de previsión: conviene fusionarlos para no '
        'contar dos veces su valor.',
        'OPP-02 cuelga de P2R-03 y de P2R-04 (aviso ⚠ en la aplicación); lo cura P2R-04, así que se puede quitar de P2R-03.',
        'En la aplicación la descripción de cada hallazgo es solo el nombre del tipo; aquí van las descripciones acordadas, '
        'para copiarlas.',
    ],
    'MGT': [
        'OPP-02 cuelga de P2R-03 y de P2R-04 (aviso ⚠ en P2R-03); lo cura P2R-04, así que se puede quitar de P2R-03.',
        'P2R-05 no cura HT-P2R-06 en MGT (no está asignado), solo HT-P2R-07.',
        'En la aplicación la descripción de cada hallazgo es solo el nombre del tipo; aquí van las descripciones acordadas, '
        'para copiarlas.',
    ],
    'Ondupack Almendralejo': [
        'El tipo HT-P2R-01 es ahora Oportunidad, pero INF-01 sigue como Ineficiencia · Circulante (igual que INF-10 e INF-16 '
        'en Ondupet).',
        'ACC-07 es un agente que no cambia ninguna caja: si se marca sobre «Autorizar la liberación con material pendiente», '
        'aparecerá en el to-be y tendrá capacidad.',
        'HT-P2R-12 no está asignado en esta celda, así que P2R-08 solo cura OPP-01.',
    ],
}

THIN = Side(style='thin', color=BORDER)
GRUESO = Side(style='medium', color=NAVY)


def hoja_empresa(wb, empresa, palancas_cat):
    """Una hoja por empresa, organizada por palanca: hallazgos que cura e intervenciones, lado a lado."""
    ws = wb.create_sheet(empresa)
    hall = HALLAZGOS[empresa]
    interv = INTERVENCIONES[empresa]
    propuesta = empresa == 'MGT'
    titulo(ws, f'P2R × {empresa} · hallazgos, palancas e intervenciones', CABECERA[empresa])
    cab = ['Palanca', 'Qué cambia', 'Hallazgo', 'Familia · naturaleza', 'Anclaje (actividad · fase)', 'Descripción del hallazgo',
           'Intervención', 'De qué es', 'Qué hay que hacer', 'Caja que cambia', 'Capacidad (est.)']
    fila_cab = 6
    cabecera(ws, fila_cab, cab, [24, 40, 26, 18, 26, 52, 32, 10, 58, 20, 11])
    r = fila_cab + 1
    for pid in PALANCAS[empresa]:
        nombre, que = palancas_cat[pid]
        hs = [h for h in hall if pid in h[5]]
        its = [i for i in interv if i[3] == pid]
        n = max(len(hs), len(its), 1)
        for k in range(n):
            fila = r + k
            if k < len(hs):
                hid, tipo, famnat, anclaje, desc, pals = hs[k]
                otras = [p for p in pals if p != pid]
                act, fase = anclaje.split(' · ', 1)
                vals = [f'{hid}\n{tipo}', famnat, act if act == fase else anclaje, desc + (f' (también en {", ".join(otras)})' if otras else '')]
                for j, v in enumerate(vals, 3):
                    celda(ws, fila, j, v, centrada=j == 4)
            else:
                for j in range(3, 7):
                    celda(ws, fila, j, None)
            if k < len(its):
                cod, nom, tipo, _, caja, cap, desc = its[k]
                vals = [f'{cod} · {nom}' if not propuesta else nom, tipo, descripcion(empresa, desc), caja, cap]
                for j, v in enumerate(vals, 7):
                    c = celda(ws, fila, j, v, centrada=j in (8, 10, 11), fmt='#,##0 "€"' if j == 11 else None)
                    if propuesta:
                        c.font = Font(name=F, size=10, italic=True, color=GRIS)
                    if j == 11 and v is None:
                        c.value = '—'
            else:
                for j in range(7, 12):
                    celda(ws, fila, j, None)
        for j, v in ((1, f'{pid}\n{nombre}'), (2, que)):
            c = celda(ws, r, j, v, bold=j == 1)
            c.fill = fill(TINT)
            c.alignment = Alignment(wrap_text=True, vertical='top')
            if n > 1:
                ws.merge_cells(start_row=r, start_column=j, end_row=r + n - 1, end_column=j)
        for j in range(1, 12):
            c = ws.cell(r, j)
            c.border = Border(left=THIN, right=THIN, top=GRUESO, bottom=THIN)
        r += n
    ultima = r - 1
    # resumen con fórmulas sobre la tabla
    rng = lambda col: f'{col}{fila_cab + 1}:{col}{ultima}'
    resumen = [('Hallazgos', f'=SUMPRODUCT(({rng("C")}<>"")/COUNTIF({rng("C")},{rng("C")}&""))'),
               ('Palancas', f'=COUNTA({rng("A")})'),
               ('Intervenciones', f'=COUNTA({rng("G")})'),
               ('De ellas, agentes', f'=COUNTIF({rng("H")},"agente")'),
               ('Capacidad (est.)', f'=SUM({rng("K")})')]
    for k, (etq, formula) in enumerate(resumen):
        col = 1 + 2 * k
        a = ws.cell(4, col, etq)
        a.font = Font(name=F, size=9, color=GRIS)
        b = ws.cell(4, col + 1, formula)
        b.font = Font(name=F, size=11, bold=True, color=NAVY)
        b.alignment = Alignment(horizontal='left')
        if etq.startswith('Capacidad'):
            b.number_format = '#,##0 "€"'
    nr = ultima + 2
    ws.cell(nr, 1, 'Notas').font = Font(name=F, size=11, bold=True, color=NAVY)
    for k, nota in enumerate(NOTAS[empresa], 1):
        c = ws.cell(nr + k, 1, f'· {nota}')
        c.font = Font(name=F, size=10, color=NAVY)
    ws.freeze_panes = ws.cell(fila_cab + 1, 3)
    ws.sheet_view.showGridLines = False
    return ws
